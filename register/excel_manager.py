import os
import datetime
import json
import openpyxl
import yaml
import shutil
from pathlib import Path


class ExcelManager:
    def __init__(self):
        os.system("cls")

        # init 파일 경로 설정
        pdir = Path(__file__).parent
        meta_path = pdir / ".." / "bin" / "settings" / "meta.yaml"

        # settings 경로 체크
        if not os.path.exists(meta_path.parent):
            print("[ERR] meta.yaml 파일을 읽어들일 수 없습니다.")
            return

        # 파일 읽어서 meta 객체 초기화
        with open(meta_path, "r", encoding="utf8") as fp:
            self.meta = yaml.load(fp, Loader = yaml.FullLoader)

        self._path = self.meta['path']
        self._original_file_path = pdir / ".." / "bin" / "settings" / "patch.xlsx"
        self._data_file_path = pdir / Path(self._path['data'])

        # 원본 Excel 파일이 있는지 검사
        if not os.path.exists(self._original_file_path):
            print("원본 파일이 존재하지 않습니다.")
            print("프로그램을 종료합니다...")
            return

        if not (self._data_file_path / "result.json").exists():
            print("result.json 파일이 존재하지 않습니다.")
            print("프로그램을 종료합니다...")
            return
        
        with open(self._data_file_path / "result.json", "r", encoding="utf8") as fp:
            self.result_dict = json.load(fp)
            print("JSON 파일 로딩 완료")

        default_dir = Path.home() / "Desktop" / "dotnet_excel"

        if not os.path.exists(default_dir):
            print("복사 파일이 저장될 폴더를 만듭니다.")
            os.mkdir(default_dir)

        now = datetime.datetime.now().strftime("%Y%m%d")
        self.excel_file_path = default_dir / f"patch{now}.xlsx"

        shutil.copy(self._original_file_path, self.excel_file_path)


        sheet = input("어떤 Sheet를 선택할까요? (dotnet/3rdparty) ")
        
        if sheet == '3rdparty':
            self.patch = input("어떤 패치를 입력할까요? (java/adobe) ")

            if self.patch == "adobe":
                self.adobe_version = input("(Continuos, Classic)")

        self.excel_file = openpyxl.load_workbook(self.excel_file_path)
        self.excel_sheet = self.excel_file.get_sheet_by_name(sheet)
        print("엑셀 파일 로딩이 완료되었습니다...")
    

    def get_cell_value(self, row, col: str):
        return self.excel_sheet.cell(row, self._col_to_int(col)).value


    def set_cell_value(self, row, col: str, val):
        self.excel_sheet.cell(row, self._col_to_int(col), val)


    def save_workbook(self):
        print("Excel 작업을 마무리합니다...")
        self.excel_file.save(self.excel_file_path)


    def _col_to_int(self, col: str) -> int:
        val = ord(col) - ord('A') if col.isupper() else ord(col) - ord('a')
        return val + 1
    

    def _work_sheet_props(self):
        sheet = self.excel_sheet
        print(sheet.sheet_properties)
    

if __name__ == "__main__":
    em = ExcelManager()
    # em.save_workbook()
    # em._work_sheet_props()